"""Results PowerPoint deck — a pre-populated, Castalia-style presentation built from the live engine run.

Structure: title → executive summary (coverage table + headline text) → per sector two slides (a
safely-managed coverage slide with the stacked chart + written summary, and a financing-gap slide with the
gap chart + the per-intervention contribution table). Charts are captured on the client (recharts SVG →
canvas) and passed in as data-URL PNGs; if absent the slides still render the tables and text.
"""

import io
import base64
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

from export_data import intervention_breakdown, WATER_INTV, SAN_INTV, _cur

# ── palette (Castalia dark-blue) ────────────────────────────────────────────────────────────────────
NAVY = RGBColor(0x0B, 0x25, 0x45)
NAVY2 = RGBColor(0x13, 0x34, 0x5F)
SKY = RGBColor(0x0E, 0xA5, 0xE9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1E, 0x3A, 0x5F)
GREY = RGBColor(0x64, 0x74, 0x8B)
ROW = RGBColor(0xF1, 0xF8, 0xFD)
BAU_C = RGBColor(0x25, 0x63, 0xEB)
TGT_C = RGBColor(0x16, 0xA3, 0x4A)
SCN_C = RGBColor(0xEA, 0x58, 0x0C)


def _pct(f):
    return f"{f * 100:.1f}%"


def _b(v_millions):
    return f"{v_millions / 1000:,.1f}"


def _sector_summary(result, inputs, sk):
    years = result['years']
    per = inputs.get('period', {})
    by = per.get('baseline_year', years[0])
    e = len(years) - 1
    bi = max(0, years.index(by)) if by in years else 0
    sec = result[sk]
    total = result['total_hh']

    def cov(arr, i):
        t = total[i] or 0
        return (min(t, (arr[i] or 0)) / t) if t else 0.0
    bau, scn, tgt = sec['bau_hh'][0], sec['scenario_hh'][0], sec['target_hh'][0]
    cum = lambda a: sum((a[i] or 0) for i, y in enumerate(years) if y > by)
    return {
        'end': years[e], 'curCov': cov(bau, bi), 'bauCov': cov(bau, e), 'scnCov': cov(scn, e), 'tgtCov': cov(tgt, e),
        'addHH': max(0.0, (min(total[e], scn[e]) - min(total[e], bau[e]))),
        'gapBau': cum(sec.get('financing_gap') or []), 'gapScn': cum(sec.get('scenario_financing_gap') or []),
    }


def create_pptx(result: dict, inputs: dict, charts: dict | None = None) -> io.BytesIO:
    charts = charts or {}
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    cc = inputs.get('country_config', {})
    period = inputs.get('period', {})
    country = cc.get('country', 'Country')
    area = cc.get('area', 'Area')
    cur = _cur(inputs)
    baseline = period.get('baseline_year', 2025)
    end = result['years'][-1]

    def textbox(slide, l, t, w, h, anchor=None):
        tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        if anchor is not None:
            tf.vertical_anchor = anchor
        return tf

    def set_p(p, text, size, color, bold=False, bullet=False):
        p.text = ('• ' + text) if bullet else text
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.font.bold = bold

    def band(slide, color, top, height):
        from pptx.enum.shapes import MSO_SHAPE
        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(top), prs.slide_width, Inches(height))
        shp.fill.solid(); shp.fill.fore_color.rgb = color; shp.line.fill.background()
        shp.shadow.inherit = False
        return shp

    def slide_title(slide, title, subtitle=None):
        band(slide, NAVY, 0, 1.0)
        tf = textbox(slide, 0.5, 0.12, 12.3, 0.8, MSO_ANCHOR.MIDDLE)
        set_p(tf.paragraphs[0], title, 22, WHITE, bold=True)
        if subtitle:
            p = tf.add_paragraph(); set_p(p, subtitle, 12, RGBColor(0xBF, 0xDB, 0xFE))

    def add_table(slide, left, top, width, headers, rows, col0_left=True, fontsize=11, header_fill=SKY, total_last=False):
        n_rows, n_cols = len(rows) + 1, len(headers)
        tbl = slide.shapes.add_table(n_rows, n_cols, Inches(left), Inches(top), Inches(width), Inches(0.34 * n_rows)).table
        for j, h in enumerate(headers):
            c = tbl.cell(0, j); c.text = str(h)
            c.fill.solid(); c.fill.fore_color.rgb = header_fill
            for p in c.text_frame.paragraphs:
                p.font.size = Pt(fontsize); p.font.bold = True; p.font.color.rgb = WHITE
                p.alignment = PP_ALIGN.LEFT if (j == 0 and col0_left) else PP_ALIGN.RIGHT
        for i, row in enumerate(rows):
            is_total = total_last and i == len(rows) - 1
            for j, val in enumerate(row):
                c = tbl.cell(i + 1, j); c.text = str(val)
                c.fill.solid(); c.fill.fore_color.rgb = (RGBColor(0xDF, 0xF1, 0xFB) if is_total else (ROW if i % 2 else WHITE))
                for p in c.text_frame.paragraphs:
                    p.font.size = Pt(fontsize - 0.5); p.font.bold = is_total
                    p.font.color.rgb = INK if is_total else RGBColor(0x33, 0x41, 0x55)
                    p.alignment = PP_ALIGN.LEFT if (j == 0 and col0_left) else PP_ALIGN.RIGHT
        return tbl

    def add_chart(slide, key, left, top, width):
        data_url = charts.get(key)
        if not data_url or ',' not in data_url:
            tf = textbox(slide, left, top + 1.6, width, 0.6, MSO_ANCHOR.MIDDLE)
            set_p(tf.paragraphs[0], '(chart unavailable)', 11, GREY)
            return
        raw = base64.b64decode(data_url.split(',', 1)[1])
        slide.shapes.add_picture(io.BytesIO(raw), Inches(left), Inches(top), width=Inches(width))

    # === 1. TITLE ===
    s = prs.slides.add_slide(blank)
    bg = s.background.fill; bg.solid(); bg.fore_color.rgb = NAVY
    tf = textbox(s, 1.0, 2.5, 11.3, 2.4)
    set_p(tf.paragraphs[0], 'Water & Sanitation — Strategic Scenarios', 40, WHITE, bold=True)
    p = tf.add_paragraph(); set_p(p, f'{country} · {area}', 20, SKY)
    p = tf.add_paragraph(); set_p(p, f'Baseline {baseline} · forecast to {end}', 14, RGBColor(0x94, 0xA3, 0xB8))
    band(s, SKY, 7.15, 0.35)

    # === 2. EXECUTIVE SUMMARY ===
    s = prs.slides.add_slide(blank)
    slide_title(s, 'Executive summary', f'Safely-managed coverage and financing gap · {country} {area}')
    ws, sn = _sector_summary(result, inputs, 'water_supply'), _sector_summary(result, inputs, 'sanitation')
    headers = ['Sector', 'Current', f'BAU {end}', f'Target {end}', f'With reforms {end}']
    rows = [
        ['Water Supply', _pct(ws['curCov']), _pct(ws['bauCov']), _pct(ws['tgtCov']), _pct(ws['scnCov'])],
        ['Sanitation', _pct(sn['curCov']), _pct(sn['bauCov']), _pct(sn['tgtCov']), _pct(sn['scnCov'])],
    ]
    add_table(s, 0.6, 1.4, 8.0, headers, rows)
    tf = textbox(s, 0.6, 3.2, 12.1, 3.6)
    set_p(tf.paragraphs[0], 'Key takeaways', 15, INK, bold=True)
    for label, d in [('Water supply', ws), ('Sanitation', sn)]:
        red = (1 - d['gapScn'] / d['gapBau']) * 100 if d['gapBau'] else 0
        txt = (f"{label}: safely-managed coverage reaches {_pct(d['scnCov'])} with the current interventions by "
               f"{d['end']} (vs {_pct(d['bauCov'])} business-as-usual and a {_pct(d['tgtCov'])} target) — "
               f"{d['addHH']:.2f} M more households. Cumulative financing gap falls from {_b(d['gapBau'])} to "
               f"{_b(d['gapScn'])} B {cur} ({red:.0f}% lower).")
        p = tf.add_paragraph(); set_p(p, txt, 12.5, RGBColor(0x33, 0x41, 0x55), bullet=True)

    # === 3. PER-SECTOR SLIDES ===
    for sk, name, defs, cov_key, gap_key in [
        ('water_supply', 'Water Supply', WATER_INTV, 'water_coverage', 'water_gap'),
        ('sanitation', 'Sanitation', SAN_INTV, 'san_coverage', 'san_gap'),
    ]:
        d = _sector_summary(result, inputs, sk)

        # 3a. coverage chart + written summary
        s = prs.slides.add_slide(blank)
        slide_title(s, f'{name} — safely-managed coverage', 'BAU baseline plus each intervention’s contribution')
        add_chart(s, cov_key, 0.5, 1.35, 7.6)
        tf = textbox(s, 8.4, 1.5, 4.5, 5.4)
        set_p(tf.paragraphs[0], f'By {d["end"]}', 15, INK, bold=True)
        bullets = [
            f'BAU coverage: {_pct(d["bauCov"])}',
            f'With interventions: {_pct(d["scnCov"])}',
            f'Target: {_pct(d["tgtCov"])}',
            f'Extra safely-managed: {d["addHH"]:.2f} M households',
        ]
        for bt in bullets:
            p = tf.add_paragraph(); set_p(p, bt, 13, RGBColor(0x33, 0x41, 0x55), bullet=True)

        # 3b. financing-gap chart + per-intervention contribution table
        s = prs.slides.add_slide(blank)
        slide_title(s, f'{name} — financing gap & interventions', 'Grey = gap remaining · colours = closed by each lever')
        add_chart(s, gap_key, 0.5, 1.35, 7.6)
        bd = intervention_breakdown(inputs, sk, defs)
        headers = ['Intervention', f'Added HH (M)', f'Resources ({cur} B)', f'Gap closed ({cur} B)']
        if bd:
            rows = [[lbl, f'{hh:.3f}', ('—' if res is None else f'{res:,.2f}'), f'{gap:,.2f}'] for (lbl, hh, res, gap) in bd]
            rows.append(['Total', f'{sum(x[1] for x in bd):.3f}', f'{sum((x[2] or 0) for x in bd):,.2f}', f'{sum(x[3] for x in bd):,.2f}'])
            add_table(s, 8.3, 1.5, 4.7, headers, rows, fontsize=10, total_last=True)
        else:
            tf = textbox(s, 8.3, 2.6, 4.7, 1.0)
            set_p(tf.paragraphs[0], 'No interventions enabled for this sector.', 12, GREY)
        tf = textbox(s, 8.3, 6.2, 4.7, 1.0)
        red = (1 - d['gapScn'] / d['gapBau']) * 100 if d['gapBau'] else 0
        set_p(tf.paragraphs[0], f'Cumulative gap: {_b(d["gapBau"])} → {_b(d["gapScn"])} B {cur} ({red:.0f}% lower).', 11.5, INK, bold=True)

    output = io.BytesIO()
    prs.save(output)
    output.seek(0)
    return output
